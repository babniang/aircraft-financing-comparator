"""Formula-driven Excel model export.

Produces a multi-tab .xlsx that mirrors ``financial_model.py`` as live Excel
formulas, so a reader can open it and flex every input without the API:

    Summary        cover page: the method in one line, the cheapest structure
                   with its execution caveat beside it, the four implied
                   costs, deal inputs, data provenance and a tab guide
    Assumptions    every input in one place (blue = input, black = formula,
                   green = linked from another tab)
    ECA Debt, SLL, JOLCO, SLB
                   key terms, the full cash-flow schedule, the IRR, formula
                   integrity checks and the credit metrics, all live
    Comparison     the structures side by side, credit metrics, execution
                   notes and a chart

Nothing in the workbook is a pasted model output: every result is a formula a
reviewer can click into. Agreement with the Python model is enforced by the
test suite (tests/test_api.py recalculates the workbook and compares), not
asserted in a cell.
"""

from __future__ import annotations

import datetime as _dt
from io import BytesIO
from typing import Callable

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.axis import ChartLines
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils.indexed_list import IndexedList
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

import financial_model as fm

# --- house style (matches the web app: white, near-black, one blue) -------

FONT = "Calibri"
BASE = 12

C_INK = "111111"
C_MUTED = "6E777A"
C_BLUE = "2A6BF2"
C_BLUE_DARK = "1B4BC0"
C_BLUE_PALE = "F2F6FF"
C_BLUE_TINT = "E3EDFF"
C_RULE = "DFE3E6"
C_HEAD = "1B2430"
C_GREEN = "0B7A3E"
C_RED = "C62828"
C_AMBER = "7A4E00"


def _font(size: float = BASE, bold: bool = False, italic: bool = False,
          color: str = C_INK) -> Font:
    return Font(name=FONT, size=size, bold=bold, italic=italic, color=color)


F_TITLE = _font(22, bold=True)
F_SUB = _font(12, color=C_MUTED)
F_SECTION = _font(12, bold=True, color="FFFFFF")
F_HEAD = _font(11, bold=True, color="FFFFFF")
F_LABEL = _font()
F_LABEL_B = _font(bold=True)
F_INPUT = _font(bold=True, color=C_BLUE_DARK)
F_CALC = _font()
F_CALC_B = _font(bold=True)
F_LINK = _font(color=C_GREEN)
F_LINK_B = _font(bold=True, color=C_GREEN)
F_LINK_KPI = _font(18, bold=True, color=C_GREEN)
F_NOTE = _font(11, italic=True, color=C_MUTED)
F_KPI = _font(16, bold=True, color=C_BLUE_DARK)
F_OK = _font(bold=True, color=C_GREEN)
F_CAVEAT = _font(12, color=C_AMBER)
F_CAVEAT_B = _font(12, bold=True, color=C_AMBER)

FILL_SECTION = PatternFill("solid", fgColor=C_BLUE)
FILL_HEAD = PatternFill("solid", fgColor=C_HEAD)
FILL_INPUT = PatternFill("solid", fgColor=C_BLUE_PALE)
FILL_LEAD = PatternFill("solid", fgColor=C_BLUE_TINT)
FILL_BOX = PatternFill("solid", fgColor="F5F7FA")
FILL_CAVEAT = PatternFill("solid", fgColor="FFF4DE")
FILL_TOTAL = PatternFill("solid", fgColor="F3F5F7")

_THIN = Side(style="thin", color=C_RULE)
_INK_LINE = Side(style="thin", color=C_INK)
B_ROW = Border(bottom=_THIN)
B_BOX = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
B_TOTAL = Border(top=_INK_LINE)

NUM_USD = '#,##0.00_);[Red](#,##0.00);"-"_)'
NUM_PCT = "0.00%"
NUM_PCT1 = "0.0%"
NUM_RATE = "0.00"
NUM_INT = "0"
NUM_BPS0 = '#,##0" bps"'
NUM_BPS = '#,##0.0" bps";[Red]-#,##0.0" bps"'
NUM_EDGE = '+#,##0.0" bps";-#,##0.0" bps";"lowest"'
NUM_X = '0.00"x"'
NUM_YRS = '0.0" yrs"'
NUM_CHECK = '0.000000;[Red]-0.000000;0.000000'

# Integrity-check tolerance, in $m: $1k. Excel's IRR converges to ~1e-7 in
# rate, which moves the NPV of a ~$50m schedule by a few tens of dollars.
CHECK_TOL = 0.001

# Execution friction per structure, shown next to the cheapest result. The
# headline rate deliberately excludes these one-off and practical costs.
EXECUTION_NOTES = {
    "eca_debt": (
        "Needs export-credit-agency eligibility (buyer credit, country cover, OECD "
        "Aircraft Sector Understanding terms). The guarantee premium is paid upfront "
        "and documentation runs through the agency, so lead times exceed a bank loan."
    ),
    "sll": (
        "Needs a credible sustainability KPI framework (fleet CO2 intensity, SAF "
        "uptake) with annual third-party verification. The ratchet is small, so the "
        "label matters more for lender appetite than for cost."
    ),
    "jolco": (
        "Lowest headline rate, hardest to execute: a Japanese SPC, Japanese "
        "tax-equity investors with appetite that year, a strong lessee credit and a "
        "long tenor. Arrangement, equity-placement and legal costs are excluded from "
        "this rate and are largely fixed, so the edge narrows or reverses on smaller "
        "deals. Prepayment is restricted until the call."
    ),
    "slb": (
        "Fastest to execute, lessor-led. The aircraft leaves the airline's fixed "
        "assets, but IFRS 16 brings a lease liability on instead. The real cost is "
        "the residual handed to the lessor, plus return conditions at lease end."
    ),
}

ACCOUNTING = {
    "eca_debt": "Owned asset, loan on balance sheet",
    "sll": "Owned asset, loan on balance sheet",
    "jolco": "Finance lease (IFRS 16)",
    "slb": "Asset sold; lease liability (IFRS 16)",
}

ORDER = ["eca_debt", "sll", "jolco", "slb"]
SHEET = {"eca_debt": "ECA Debt", "sll": "SLL", "jolco": "JOLCO", "slb": "SLB"}

# --- low-level helpers ----------------------------------------------------


def _set(ws: Worksheet, ref: str, value, *, font: Font = F_CALC, num: str | None = None,
         fill: PatternFill | None = None, align: str | None = None,
         border: Border | None = None, wrap: bool = False, indent: int = 0):
    c = ws[ref]
    c.value = value
    c.font = font
    if num:
        c.number_format = num
    if fill:
        c.fill = fill
    if border:
        c.border = border
    c.alignment = Alignment(horizontal=align, vertical="top" if wrap else "center",
                            wrap_text=wrap, indent=indent)
    return c


def _merge(ws: Worksheet, rng: str, value, **kw):
    ws.merge_cells(rng)
    return _set(ws, rng.split(":")[0], value, **kw)


def _band(ws: Worksheet, row: int, text: str, last_col: int = 7) -> None:
    for col in range(1, last_col + 1):
        ws.cell(row=row, column=col).fill = FILL_SECTION
    c = ws.cell(row=row, column=1, value=text)
    c.font = F_SECTION
    c.alignment = Alignment(vertical="center")
    ws.row_dimensions[row].height = 24


def _header(ws: Worksheet, row: int, headers: list[str], left: tuple[int, ...] = ()) -> None:
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = F_HEAD
        c.fill = FILL_HEAD
        c.alignment = Alignment(horizontal="left" if (i == 1 or i in left) else "right",
                                vertical="center", wrap_text=True,
                                indent=1 if i in left else 0)
    ws.row_dimensions[row].height = 32


def _sheet(ws: Worksheet, widths: dict[str, float], tab: str) -> None:
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 100
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.sheet_properties.tabColor = tab
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _title(ws: Worksheet, title: str, subtitle: str) -> None:
    _set(ws, "A1", title, font=F_TITLE)
    ws.row_dimensions[1].height = 36
    _set(ws, "A2", subtitle, font=F_SUB)
    ws.row_dimensions[2].height = 20


def _set_default_font(wb: Workbook) -> None:
    """Make Calibri 12 the workbook default, so any cell a reviewer types into
    later matches the model rather than falling back to Calibri 11."""
    default = Font(name=FONT, size=BASE)
    wb._fonts = IndexedList([default])
    wb._named_styles["Normal"].font = default


# --- Assumptions tab: the single source of inputs ------------------------

A = "Assumptions"
REF: dict[str, str] = {}


def R(key: str) -> str:
    return REF[key]


def _rate_source_text(source: str) -> str:
    return {
        "live_sofr": "live SOFR (Federal Reserve Bank of New York) at export",
        "assumption_override": "user override",
        "assumption_default": "default assumption (live SOFR unavailable or off)",
    }.get(source, source.replace("_", " "))


def _build_assumptions(ws: Worksheet, *, aircraft_label: str, delivery: float,
                       market_value: float, tenor: int, reference_rate_pct: float,
                       reference_rate_source: str, gtf: bool, a: fm.Assumptions) -> None:
    _sheet(ws, {"A": 40, "B": 16, "C": 18, "D": 104}, tab=C_BLUE_DARK)
    _title(ws, "Assumptions",
           "Every input in one place. Change a blue cell and the whole workbook recalculates.")

    _set(ws, "A3", "Colour key", font=F_LABEL_B)
    _set(ws, "B3", "Input", font=F_INPUT, fill=FILL_INPUT, align="center", border=B_BOX)
    _set(ws, "C3", "Formula", font=F_CALC, align="center", border=B_BOX)
    _set(ws, "D3", "Green text on other tabs = linked from another tab", font=F_LINK)

    # (key, label, value, number format, unit, basis note, editable)
    rows: list = [
        ("SECTION", "Deal"),
        ("aircraft", "Aircraft type", aircraft_label, None, "",
         "From the tool's reference data.", False),
        ("delivery", "Delivery price", delivery, NUM_USD, "$m",
         "Acquisition cost paid for the aircraft.", True),
        ("market_value", "Market value today", market_value, NUM_USD, "$m",
         "Current value of the airframe: the base for the SLB sale price and every residual.", True),
        ("tenor", "Tenor", tenor, NUM_INT, "years",
         "Fixed: schedules are built to this length. Regenerate from the tool to change it.", False),
        ("reference_rate_pct", "Reference rate", reference_rate_pct, NUM_RATE, "%",
         f"Source: {_rate_source_text(reference_rate_source)}.", True),
        ("gtf", "GTF engine-risk adjustment", 1 if gtf else 0, NUM_INT, "1 = on, 0 = off",
         "Pratt & Whitney GTF powder-metal exposure: adds debt margin and cuts values.", True),

        ("SECTION", "ECA-backed debt"),
        ("eca_ltv_pct", "ECA loan-to-value", a.eca_ltv_pct, NUM_PCT1, "of price",
         "Share of the delivery price the guaranteed loan covers (ECA cover runs up to 85%).", True),
        ("eca_base_margin_bps", "ECA loan margin", a.eca_base_margin_bps, NUM_INT, "bps",
         "Bank margin over the reference rate; tight because the agency carries the credit risk.", True),
        ("eca_guarantee_fee_pct", "ECA guarantee premium", a.eca_guarantee_fee_pct, NUM_RATE,
         "% of loan", "Upfront, risk-based under the OECD Aircraft Sector Understanding; "
         "spread over the tenor for the headline cost.", True),

        ("SECTION", "Sustainability-linked loan"),
        ("sll_ltv_pct", "SLL loan-to-value", a.sll_ltv_pct, NUM_PCT1, "of price",
         "Commercial bank appetite without a guarantee.", True),
        ("sll_base_margin_bps", "SLL base margin", a.sll_base_margin_bps, NUM_INT, "bps",
         "Margin before the sustainability ratchet.", True),
        ("sll_kpi_ratchet_bps", "SLL KPI ratchet", a.sll_kpi_ratchet_bps, NUM_INT, "bps",
         "Margin step per KPI outcome. Base case: KPI met, margin down one step.", True),

        ("SECTION", "JOLCO (Japanese operating lease with call option)"),
        ("jolco_ltv_pct", "JOLCO financing", a.jolco_ltv_pct, NUM_PCT1, "of price",
         "Share of price funded through the Japanese SPC (senior debt plus tax equity).", True),
        ("jolco_base_margin_bps", "JOLCO senior margin", a.jolco_base_margin_bps, NUM_INT, "bps",
         "Senior lenders' margin inside the SPC; no sovereign guarantee, so above ECA.", True),
        ("jolco_tax_benefit_bps", "Japanese tax benefit passed through", a.jolco_tax_benefit_bps,
         NUM_INT, "bps", "Accelerated-depreciation benefit the equity rebates into the rentals; "
         "realistic band 50 to 150 bps.", True),
        ("jolco_call_year", "Purchase option (call) year", int(a.jolco_call_year), NUM_INT, "year",
         "Year the airline buys the aircraft back; capped at the tenor. The JOLCO tab reshapes live.", True),

        ("SECTION", "Sale-and-leaseback"),
        ("slb_lease_rate_factor_pct", "Lease rate factor", a.slb_lease_rate_factor_pct, NUM_RATE,
         "% of value / month", "IBA / Cirium convention. Annual rent = sale price x LRF x 12.", True),
        ("slb_depreciation_pct", "Value depreciation", a.slb_depreciation_pct, NUM_RATE,
         "% of today's value / yr", "Straight line, ~20-year life: the shape of airline IAS 16 policy "
         "(Lufthansa: 20 years to 5%, i.e. 4.75%/yr). Leaves ~46% at 12 years.", True),
        ("slb_residual_floor_pct", "Residual value floor", a.slb_residual_floor_pct, NUM_RATE,
         "% of today's value", "Part-out / teardown value: engines and components keep a resale value.", True),
        ("slb_residual_value_pct", "Residual override", a.slb_residual_value_pct, NUM_RATE,
         "% of today's value", "0 = use the depreciation curve above; any other value replaces it.", True),

        ("SECTION", "GTF engine risk"),
        ("gtf_margin_addon_bps", "GTF margin add-on", a.gtf_margin_addon_bps, NUM_INT, "bps",
         "Extra debt margin for GTF-powered aircraft: availability and shop-visit risk.", True),
        ("gtf_value_haircut_pct", "GTF value haircut", a.gtf_value_haircut_pct, NUM_RATE, "%",
         "Cut to the SLB sale price and residual for GTF-powered aircraft.", True),
    ]

    r = 5
    for entry in rows:
        if entry[0] == "SECTION":
            _band(ws, r, entry[1], last_col=4)
            r += 1
            continue
        key, label, value, num, unit, note, editable = entry
        _set(ws, f"A{r}", label, font=F_LABEL, border=B_ROW)
        if editable:
            _set(ws, f"B{r}", value, font=F_INPUT, num=num, fill=FILL_INPUT, align="right",
                 border=B_BOX)
        else:
            _set(ws, f"B{r}", value, font=F_CALC_B, num=num,
                 align="right" if num else "left", border=B_ROW)
        _set(ws, f"C{r}", unit, font=F_NOTE, border=B_ROW)
        _set(ws, f"D{r}", note, font=F_NOTE, border=B_ROW)
        ws.row_dimensions[r].height = 21
        REF[key] = f"'{A}'!$B${r}"
        if key == "gtf":
            dv = DataValidation(type="list", formula1='"0,1"', allow_blank=False,
                                showErrorMessage=True, error="Enter 1 (on) or 0 (off).")
            ws.add_data_validation(dv)
            dv.add(f"B{r}")
        r += 1
    ws.freeze_panes = "A5"


def _gtf_addon() -> str:
    return f"IF({R('gtf')}=1,{R('gtf_margin_addon_bps')},0)"


def _gtf_haircut() -> str:
    return f"IF({R('gtf')}=1,{R('gtf_value_haircut_pct')}/100,0)"


def _retained(years_expr: str) -> str:
    return (f"MAX({R('slb_residual_floor_pct')}/100,"
            f"1-{R('slb_depreciation_pct')}/100*{years_expr})")


# --- structure tabs -------------------------------------------------------

Term = tuple  # (key, label, builder(refs)->formula, num, explanation, font)


def _terms(ws: Worksheet, start: int, items: list[Term]) -> tuple[dict[str, str], int]:
    """Label in A:C (merged), value in D, plain-English formula in E:G.
    Builders get the refs dict so later terms can point at earlier ones."""
    refs: dict[str, str] = {}
    r = start
    for key, *_ in items:
        refs[key] = f"$D${r}"
        r += 1
    r = start
    for key, label, build, num, expl, *font in items:
        _merge(ws, f"A{r}:C{r}", label, font=F_LABEL, border=B_ROW)
        f = build(refs) if callable(build) else build
        _set(ws, f"D{r}", f, font=(font[0] if font else F_CALC_B), num=num, align="right",
             border=B_ROW)
        _merge(ws, f"E{r}:G{r}", expl, font=F_NOTE, border=B_ROW)
        ws.row_dimensions[r].height = 21
        r += 1
    return refs, r


def _schedule(ws: Worksheet, start: int, headers: list[str], tenor: int,
              row0: dict[int, str], row_t: Callable[[int, int, int], dict[int, str]],
              sum_cols: list[int]) -> tuple[int, int, int, int]:
    """Writes header, year-0 row, years 1..tenor and a totals row.
    Returns (y0_row, y1_row, last_row, next_free_row)."""
    _header(ws, start, headers)
    y0 = start + 1
    ncol = len(headers)
    ws.cell(row=y0, column=1, value=0).font = F_CALC
    for col, f in row0.items():
        c = ws.cell(row=y0, column=col, value=f)
        c.number_format = NUM_USD
        c.font = F_CALC_B
    for y in range(1, tenor + 1):
        r = y0 + y
        ws.cell(row=r, column=1, value=y).font = F_CALC
        for col, f in row_t(y, r, y0).items():
            c = ws.cell(row=r, column=col, value=f)
            c.number_format = NUM_USD
            c.font = F_CALC
    last = y0 + tenor
    for r in range(y0, last + 1):
        for col in range(1, ncol + 1):
            ws.cell(row=r, column=col).border = B_ROW
            ws.cell(row=r, column=col).alignment = Alignment(
                horizontal="left" if col == 1 else "right")
    total = last + 1
    ws.cell(row=total, column=1, value="Total")
    for col in range(1, ncol + 1):
        c = ws.cell(row=total, column=col)
        c.font = F_CALC_B
        c.fill = FILL_TOTAL
        c.border = B_TOTAL
        if col > 1:
            c.alignment = Alignment(horizontal="right")
    for col in sum_cols:
        letter = ws.cell(row=total, column=col).column_letter
        lo = y0 if col == ncol else y0 + 1
        c = ws.cell(row=total, column=col, value=f"=SUM({letter}{lo}:{letter}{last})")
        c.number_format = NUM_USD
    ws.freeze_panes = ws[f"A{start + 1}"]
    return y0, y0 + 1, last, total + 2


def _status(npv_ref: str, *more: str) -> str:
    parts = [f"ABS({npv_ref})<{CHECK_TOL}"] + [f"ABS({m})<{CHECK_TOL}" for m in more]
    return f'=IF(AND({",".join(parts)}),"OK","CHECK")'


def _check_rules(ws: Worksheet, status_ref: str) -> None:
    cell = status_ref.replace("$", "")
    ws.conditional_formatting.add(
        cell, CellIsRule(operator="equal", formula=['"CHECK"'], font=_font(bold=True, color=C_RED)))


def _eca_sheet(ws: Worksheet, tenor: int) -> dict[str, str]:
    _sheet(ws, {"A": 9, **{c: 18 for c in "BCDEFG"}}, tab=C_HEAD)
    _title(ws, "ECA-Backed Term Loan",
           "Export-credit-agency guaranteed loan: tight margin, upfront guarantee premium, fully amortising.")
    _band(ws, 4, "Key terms")
    t, r = _terms(ws, 5, [
        ("loan", "Loan amount ($m)", f"={R('delivery')}*{R('eca_ltv_pct')}", NUM_USD,
         "Delivery price x ECA loan-to-value"),
        ("rate", "All-in interest rate",
         f"={R('reference_rate_pct')}/100+({R('eca_base_margin_bps')}+{_gtf_addon()})/10000",
         NUM_PCT, "Reference rate + ECA margin (+ GTF add-on when on)"),
        ("ds", "Annual debt service ($m)",
         lambda t: f"=-PMT({t['rate']},{R('tenor')},{t['loan']})", NUM_USD,
         "Level annuity that fully repays the loan over the tenor"),
        ("fee", "Guarantee premium, upfront ($m)",
         lambda t: f"={t['loan']}*{R('eca_guarantee_fee_pct')}/100", NUM_USD,
         "Premium x loan amount, paid to the agency at drawdown"),
        ("fee_yr", "Guarantee premium, per year ($m)",
         lambda t: f"={t['fee']}/{R('tenor')}", NUM_USD,
         "Spread straight-line over the tenor for the headline cost"),
    ])

    _band(ws, r + 1, "Cash-flow schedule ($m, incremental versus paying cash at delivery)")
    y0, y1, last, nxt = _schedule(
        ws, r + 2,
        ["Year", "Opening balance", "Interest", "Principal", "Debt service",
         "Guarantee premium", "Incremental cash flow"],
        tenor,
        row0={2: f"={t['loan']}", 7: f"={t['loan']}"},
        row_t=lambda y, r_, y0_: {
            2: f"=B{y0_}" if y == 1 else f"=B{r_ - 1}-D{r_ - 1}",
            3: f"=B{r_}*{t['rate']}",
            4: f"={t['ds']}-C{r_}",
            5: f"={t['ds']}",
            6: f"={t['fee_yr']}",
            7: f"=-(E{r_}+F{r_})",
        },
        sum_cols=[3, 4, 5, 6, 7],
    )

    _band(ws, nxt, "Result and integrity checks")
    res, r = _terms(ws, nxt + 1, [
        ("irr", "Implied annual cost (IRR)", f"=IRR(G{y0}:G{last})", NUM_PCT,
         "Discount rate at which the incremental cash flows net to zero", F_KPI),
        ("npv", "Check: NPV of the cash flows at that IRR ($m)",
         lambda x: f"=G{y0}+NPV({x['irr']},G{y1}:G{last})", NUM_CHECK,
         "Must be zero: proves the IRR solves this exact series"),
        ("prin", "Check: principal repaid less loan amount ($m)",
         f"=SUM(D{y1}:D{last})-{t['loan']}", NUM_CHECK,
         "Must be zero: the schedule fully amortises the loan"),
        ("drag", "Cost of the guarantee premium",
         lambda x: f"=({x['irr']}-{t['rate']})*10000", NUM_BPS,
         "IRR minus the loan rate: what the upfront premium adds each year"),
        ("status", "Integrity status", lambda x: _status(x["npv"], x["prin"]), None,
         "OK when every check above is within $1k", F_OK),
    ])
    _check_rules(ws, res["status"])

    _band(ws, r + 1, "Credit metrics (aircraft-level proxies)")
    mv = R("market_value")
    cm, _ = _terms(ws, r + 2, [
        ("ltv0", "Initial LTV", f"={t['loan']}/{mv}", NUM_PCT1, "Loan / market value today"),
        ("mid", "Mid-life year", f"=MAX(1,INT({R('tenor')}/2))", NUM_INT, "Half the tenor"),
        ("debt_mid", "Debt outstanding at mid-life ($m)",
         lambda x: f"=INDEX(B{y1}:B{last},{x['mid']})-INDEX(D{y1}:D{last},{x['mid']})",
         NUM_USD, "Closing balance at the end of that year"),
        ("val_mid", "Aircraft value at mid-life ($m)",
         lambda x: f"={mv}*{_retained(x['mid'])}", NUM_USD,
         "Today's value on the depreciation curve (Assumptions)"),
        ("ltv_mid", "LTV at mid-life", lambda x: f"={x['debt_mid']}/{x['val_mid']}", NUM_PCT1,
         "Is the loan amortising faster than the asset loses value?"),
        ("dscr", "Asset DSCR",
         f"=({mv}*{R('slb_lease_rate_factor_pct')}/100*12)/({t['ds']}+{t['fee_yr']})", NUM_X,
         "Aircraft's lease-earning power / annual financing outflow"),
        ("wal", "Weighted average life",
         f"=SUMPRODUCT(A{y1}:A{last},D{y1}:D{last})/{t['loan']}", NUM_YRS,
         "Principal-weighted average repayment date"),
        ("balloon", "Balloon at maturity", f"=(B{last}-D{last})/{t['loan']}", NUM_PCT1,
         "Share of the loan still due at the end: zero, fully amortising"),
    ])
    return {**res, **cm, "financed": R("eca_ltv_pct")}


def _sll_sheet(ws: Worksheet, tenor: int) -> dict[str, str]:
    _sheet(ws, {"A": 9, **{c: 18 for c in "BCDEFG"}}, tab=C_HEAD)
    _title(ws, "Sustainability-Linked Loan",
           "Commercial bank loan whose margin steps with a sustainability KPI. No guarantee, no upfront fee.")
    _band(ws, 4, "Key terms")
    t, r = _terms(ws, 5, [
        ("loan", "Loan amount ($m)", f"={R('delivery')}*{R('sll_ltv_pct')}", NUM_USD,
         "Delivery price x SLL loan-to-value"),
        ("margin", "Margin after KPI ratchet",
         f"={R('sll_base_margin_bps')}+{_gtf_addon()}-{R('sll_kpi_ratchet_bps')}", NUM_BPS0,
         "Base margin (+ GTF add-on) less one ratchet: KPI assumed met"),
        ("rate", "All-in interest rate",
         lambda t: f"={R('reference_rate_pct')}/100+{t['margin']}/10000", NUM_PCT,
         "Reference rate + margin after ratchet"),
        ("ds", "Annual debt service ($m)",
         lambda t: f"=-PMT({t['rate']},{R('tenor')},{t['loan']})", NUM_USD,
         "Level annuity over the tenor"),
        ("miss", "Margin if the KPI is missed",
         lambda t: f"={t['margin']}+2*{R('sll_kpi_ratchet_bps')}", NUM_BPS0,
         "Missing the target swings the margin up two steps"),
    ])

    _band(ws, r + 1, "Cash-flow schedule ($m, incremental versus paying cash at delivery)", last_col=7)
    y0, y1, last, nxt = _schedule(
        ws, r + 2,
        ["Year", "Opening balance", "Interest", "Principal", "Debt service",
         "Incremental cash flow"],
        tenor,
        row0={2: f"={t['loan']}", 6: f"={t['loan']}"},
        row_t=lambda y, r_, y0_: {
            2: f"=B{y0_}" if y == 1 else f"=B{r_ - 1}-D{r_ - 1}",
            3: f"=B{r_}*{t['rate']}",
            4: f"={t['ds']}-C{r_}",
            5: f"={t['ds']}",
            6: f"=-E{r_}",
        },
        sum_cols=[3, 4, 5, 6],
    )

    _band(ws, nxt, "Result and integrity checks")
    res, r = _terms(ws, nxt + 1, [
        ("irr", "Implied annual cost (IRR)", f"=IRR(F{y0}:F{last})", NUM_PCT,
         "Discount rate at which the incremental cash flows net to zero", F_KPI),
        ("npv", "Check: NPV of the cash flows at that IRR ($m)",
         lambda x: f"=F{y0}+NPV({x['irr']},F{y1}:F{last})", NUM_CHECK,
         "Must be zero: proves the IRR solves this exact series"),
        ("prin", "Check: principal repaid less loan amount ($m)",
         f"=SUM(D{y1}:D{last})-{t['loan']}", NUM_CHECK,
         "Must be zero: the schedule fully amortises the loan"),
        ("gap", "Check: IRR less loan rate",
         lambda x: f"=({x['irr']}-{t['rate']})*10000", NUM_BPS,
         "Must be zero: with no fees, the cost of the loan is its rate"),
        ("status", "Integrity status",
         lambda x: _status(x["npv"], x["prin"], f"{x['gap']}/100"), None,
         "OK when every check above is within tolerance", F_OK),
    ])
    _check_rules(ws, res["status"])

    _band(ws, r + 1, "Credit metrics (aircraft-level proxies)")
    mv = R("market_value")
    cm, _ = _terms(ws, r + 2, [
        ("ltv0", "Initial LTV", f"={t['loan']}/{mv}", NUM_PCT1, "Loan / market value today"),
        ("mid", "Mid-life year", f"=MAX(1,INT({R('tenor')}/2))", NUM_INT, "Half the tenor"),
        ("debt_mid", "Debt outstanding at mid-life ($m)",
         lambda x: f"=INDEX(B{y1}:B{last},{x['mid']})-INDEX(D{y1}:D{last},{x['mid']})",
         NUM_USD, "Closing balance at the end of that year"),
        ("val_mid", "Aircraft value at mid-life ($m)",
         lambda x: f"={mv}*{_retained(x['mid'])}", NUM_USD,
         "Today's value on the depreciation curve (Assumptions)"),
        ("ltv_mid", "LTV at mid-life", lambda x: f"={x['debt_mid']}/{x['val_mid']}", NUM_PCT1,
         "Is the loan amortising faster than the asset loses value?"),
        ("dscr", "Asset DSCR",
         f"=({mv}*{R('slb_lease_rate_factor_pct')}/100*12)/{t['ds']}", NUM_X,
         "Aircraft's lease-earning power / annual debt service"),
        ("wal", "Weighted average life",
         f"=SUMPRODUCT(A{y1}:A{last},D{y1}:D{last})/{t['loan']}", NUM_YRS,
         "Principal-weighted average repayment date"),
        ("balloon", "Balloon at maturity", f"=(B{last}-D{last})/{t['loan']}", NUM_PCT1,
         "Share of the loan still due at the end: zero, fully amortising"),
    ])
    return {**res, **cm, "financed": R("sll_ltv_pct")}


def _jolco_sheet(ws: Worksheet, tenor: int) -> dict[str, str]:
    _sheet(ws, {"A": 9, **{c: 18 for c in "BCDEFG"}}, tab=C_HEAD)
    _title(ws, "JOLCO (Japanese Tax Lease)",
           "Japanese SPC buys the aircraft and leases it back; tax equity rebates part of its "
           "depreciation benefit into the rent.")
    _band(ws, 4, "Key terms")
    t, r = _terms(ws, 5, [
        ("adv", "Advance via the SPC ($m)", f"={R('delivery')}*{R('jolco_ltv_pct')}", NUM_USD,
         "Delivery price x JOLCO financing share"),
        ("gross", "Senior rate before tax benefit",
         f"={R('reference_rate_pct')}/100+({R('jolco_base_margin_bps')}+{_gtf_addon()})/10000",
         NUM_PCT, "Reference rate + senior margin (+ GTF add-on when on)"),
        ("benefit", "Japanese tax benefit passed through", f"={R('jolco_tax_benefit_bps')}",
         NUM_BPS0, "Linked from Assumptions", F_LINK_B),
        ("eff", "Effective rate after tax benefit",
         lambda t: f"={t['gross']}-{t['benefit']}/10000", NUM_PCT,
         "Why JOLCO prices through plain senior debt"),
        ("rent", "Annual rental ($m)", lambda t: f"=-PMT({t['eff']},{R('tenor')},{t['adv']})",
         NUM_USD, "Level rent sized on the full tenor"),
        ("call", "Purchase option (call) year",
         f"=MIN(MAX(1,ROUND({R('jolco_call_year')},0)),{R('tenor')})", NUM_INT,
         "From Assumptions, capped at the tenor; the schedule below follows it"),
    ])

    _band(ws, r + 1, "Cash-flow schedule ($m, incremental versus paying cash at delivery)")
    call = t["call"]
    y0, y1, last, nxt = _schedule(
        ws, r + 2,
        ["Year", "Opening balance", "Interest", "Principal", "Rental", "Call price",
         "Incremental cash flow"],
        tenor,
        row0={2: f"={t['adv']}", 7: f"={t['adv']}"},
        row_t=lambda y, r_, y0_: {
            2: f"=B{y0_}" if y == 1 else f"=IF(A{r_}<={call},B{r_ - 1}-D{r_ - 1},0)",
            3: f"=IF(A{r_}<={call},B{r_}*{t['eff']},0)",
            4: f"=IF(A{r_}<={call},{t['rent']}-C{r_},0)",
            5: f"=IF(A{r_}<={call},{t['rent']},0)",
            6: f"=IF(A{r_}={call},B{r_}-D{r_},0)",
            7: f"=-(E{r_}+F{r_})",
        },
        sum_cols=[3, 4, 5, 6, 7],
    )

    _band(ws, nxt, "Result and integrity checks")
    res, r = _terms(ws, nxt + 1, [
        ("irr", "Implied annual cost (IRR)", f"=IRR(G{y0}:G{last})", NUM_PCT,
         "Discount rate at which the incremental cash flows net to zero", F_KPI),
        ("npv", "Check: NPV of the cash flows at that IRR ($m)",
         lambda x: f"=G{y0}+NPV({x['irr']},G{y1}:G{last})", NUM_CHECK,
         "Must be zero: proves the IRR solves this exact series"),
        ("prin", "Check: principal + call price less advance ($m)",
         f"=SUM(D{y1}:D{last})+SUM(F{y1}:F{last})-{t['adv']}", NUM_CHECK,
         "Must be zero: rentals and the call fully repay the SPC"),
        ("gap", "Check: IRR less effective rate",
         lambda x: f"=({x['irr']}-{t['eff']})*10000", NUM_BPS,
         "Must be zero: the call settles the balance at par"),
        ("status", "Integrity status",
         lambda x: _status(x["npv"], x["prin"], f"{x['gap']}/100"), None,
         "OK when every check above is within tolerance", F_OK),
    ])
    _check_rules(ws, res["status"])

    _band(ws, r + 1, "Credit metrics (aircraft-level proxies)")
    mv = R("market_value")
    cm, _ = _terms(ws, r + 2, [
        ("ltv0", "Initial LTV", f"={t['adv']}/{mv}", NUM_PCT1, "Advance / market value today"),
        ("mid", "Mid-life year", f"=MAX(1,INT({call}/2))", NUM_INT, "Half-way to the call"),
        ("debt_mid", "Debt outstanding at mid-life ($m)",
         lambda x: f"=INDEX(B{y1}:B{last},{x['mid']})-INDEX(D{y1}:D{last},{x['mid']})",
         NUM_USD, "Closing balance at the end of that year"),
        ("val_mid", "Aircraft value at mid-life ($m)",
         lambda x: f"={mv}*{_retained(x['mid'])}", NUM_USD,
         "Today's value on the depreciation curve (Assumptions)"),
        ("ltv_mid", "LTV at mid-life", lambda x: f"={x['debt_mid']}/{x['val_mid']}", NUM_PCT1,
         "Is the SPC amortising faster than the asset loses value?"),
        ("dscr", "Asset DSCR",
         f"=({mv}*{R('slb_lease_rate_factor_pct')}/100*12)/{t['rent']}", NUM_X,
         "Aircraft's lease-earning power / annual rental"),
        ("wal", "Weighted average life",
         f"=(SUMPRODUCT(A{y1}:A{last},D{y1}:D{last})+SUMPRODUCT(A{y1}:A{last},F{y1}:F{last}))"
         f"/{t['adv']}", NUM_YRS, "Repayments weighted by date, the call price included"),
        ("balloon", "Balloon at the call", f"=SUM(F{y1}:F{last})/{t['adv']}", NUM_PCT1,
         "Call price as a share of the advance"),
    ])
    return {**res, **cm, "financed": R("jolco_ltv_pct")}


def _slb_sheet(ws: Worksheet, tenor: int) -> dict[str, str]:
    _sheet(ws, {"A": 9, **{c: 18 for c in "BCDEFG"}}, tab=C_HEAD)
    _title(ws, "Sale-and-Leaseback",
           "The airline sells the aircraft to a lessor at market value and rents it back; "
           "the lessor keeps the residual.")
    mv = R("market_value")
    _band(ws, 4, "Key terms")
    t, r = _terms(ws, 5, [
        ("sale", "Sale price ($m)", f"={mv}*(1-{_gtf_haircut()})", NUM_USD,
         "Market value today (less GTF haircut when on)"),
        ("lrf", "Lease rate factor (% of value / month)", f"={R('slb_lease_rate_factor_pct')}",
         NUM_RATE, "Linked from Assumptions", F_LINK_B),
        ("rent", "Annual lease rental ($m)", lambda t: f"={t['sale']}*{t['lrf']}/100*12",
         NUM_USD, "Sale price x LRF x 12 (flat; a real deal often steps down)"),
        ("ret", "Value retained at lease end", f"={_retained(R('tenor'))}", NUM_PCT1,
         "Depreciation curve and floor from Assumptions, over the tenor"),
        ("resid", "Residual handed to the lessor ($m)",
         lambda t: (f"=IF({R('slb_residual_value_pct')}>0,{mv}*{R('slb_residual_value_pct')}/100,"
                    f"{mv}*{t['ret']})*(1-{_gtf_haircut()})"), NUM_USD,
         "Projected from TODAY'S value, never the original price"),
        ("gain", "Sale price less delivery price ($m)",
         lambda t: f"={t['sale']}-{R('delivery')}", NUM_USD,
         "Economic gain / (loss); IFRS 16 books only the share for rights transferred"),
    ])

    _band(ws, r + 1, "Cash-flow schedule ($m, incremental versus paying cash at delivery)")
    y0, y1, last, nxt = _schedule(
        ws, r + 2,
        ["Year", "Lease rental", "Residual given up", "Incremental cash flow"],
        tenor,
        row0={4: f"={t['sale']}"},
        row_t=lambda y, r_, y0_: {
            2: f"=-{t['rent']}",
            3: f"=IF(A{r_}={R('tenor')},-{t['resid']},0)",
            4: f"=B{r_}+C{r_}",
        },
        sum_cols=[2, 3, 4],
    )

    _band(ws, nxt, "Result and integrity checks")
    res, r = _terms(ws, nxt + 1, [
        ("irr", "Implied annual cost (IRR)", f"=IRR(D{y0}:D{last})", NUM_PCT,
         "Discount rate at which the incremental cash flows net to zero", F_KPI),
        ("npv", "Check: NPV of the cash flows at that IRR ($m)",
         lambda x: f"=D{y0}+NPV({x['irr']},D{y1}:D{last})", NUM_CHECK,
         "Must be zero: proves the IRR solves this exact series"),
        ("yield", "Rental yield on sale price", lambda x: f"={t['rent']}/{t['sale']}", NUM_PCT,
         "LRF x 12: the lessor's gross yield before residual"),
        ("status", "Integrity status", lambda x: _status(x["npv"]), None,
         "OK when the check above is within $1k", F_OK),
    ])
    _check_rules(ws, res["status"])

    _band(ws, r + 1, "Credit metrics (aircraft-level proxies)")
    cm, _ = _terms(ws, r + 2, [
        ("ltv0", "Proceeds / market value", f"={t['sale']}/{mv}", NUM_PCT1,
         "100% unless the GTF haircut applies"),
        ("ltv_mid", "LTV at mid-life", "n/a", None, "No debt balance runs down in a lease"),
        ("dscr", "Asset DSCR", f"=({mv}*{t['lrf']}/100*12)/{t['rent']}", NUM_X,
         "~1.0x by construction: rent IS the asset's earning power, so no cushion"),
        ("wal", "Weighted average life",
         f"=SUMPRODUCT(A{y1}:A{last},D{y1}:D{last})/SUM(D{y1}:D{last})", NUM_YRS,
         "Outflows (rent and residual) weighted by date"),
        ("balloon", "Balloon", "n/a", None, "Nothing to refinance; the residual is the cost"),
    ])
    return {**res, **cm, "financed": None}


# --- Comparison and Summary ----------------------------------------------


def _build_comparison(ws: Worksheet, labels: dict[str, str], refs: dict[str, dict]) -> dict:
    _sheet(ws, {"A": 34, "B": 17, "C": 16, "D": 14, "E": 34, "F": 15}, tab=C_BLUE)
    _title(ws, "Comparison", "The four structures on one basis. Every number links back to its own tab.")

    def link(key: str, field: str) -> str:
        return f"='{SHEET[key]}'!{refs[key][field]}"

    _header(ws, 3, ["Structure", "Implied annual cost", "vs lowest", "Financed",
                    "Accounting (IFRS 16)", "Integrity"], left=(5,))
    first, lastr = 4, 3 + len(ORDER)
    for i, key in enumerate(ORDER):
        r = first + i
        _set(ws, f"A{r}", labels[key], font=F_LABEL_B, border=B_ROW)
        _set(ws, f"B{r}", link(key, "irr"), font=F_LINK_B, num=NUM_PCT, align="right", border=B_ROW)
        _set(ws, f"C{r}", f"=(B{r}-MIN($B${first}:$B${lastr}))*10000", font=F_CALC,
             num=NUM_EDGE, align="right", border=B_ROW)
        fin = refs[key]["financed"]
        _set(ws, f"D{r}", f"={fin}" if fin else 1, font=F_LINK if fin else F_CALC,
             num=NUM_PCT1, align="right", border=B_ROW)
        _set(ws, f"E{r}", ACCOUNTING[key], font=F_CALC, align="left", border=B_ROW, indent=1)
        _set(ws, f"F{r}", link(key, "status"), font=F_OK, align="right", border=B_ROW)
        ws.row_dimensions[r].height = 22
        ws.conditional_formatting.add(
            f"F{r}", CellIsRule(operator="equal", formula=['"CHECK"'],
                                font=_font(bold=True, color=C_RED)))
    rng = f"$B${first}:$B${lastr}"

    _set(ws, "A10", "Lowest all-in cost", font=F_LABEL_B)
    _merge(ws, "B10:F10", f"=INDEX($A${first}:$A${lastr},MATCH(MIN({rng}),{rng},0))", font=F_KPI)
    ws.row_dimensions[10].height = 26
    _set(ws, "A11", "Edge over the next cheapest", font=F_LABEL)
    _set(ws, "B11", f"=(SMALL({rng},2)-MIN({rng}))*10000", font=F_CALC_B, num=NUM_BPS, align="right")
    _merge(ws, "C11:F11", "a year, before the one-off execution costs below", font=F_NOTE)
    _set(ws, "A12", "Worth on this deal ($m / year)", font=F_LABEL)
    _set(ws, "B12",
         f"=(SMALL({rng},2)-MIN({rng}))*INDEX($D${first}:$D${lastr},MATCH(MIN({rng}),{rng},0))"
         f"*{R('delivery')}", font=F_CALC_B, num=NUM_USD, align="right")
    _merge(ws, "C12:F12", "edge x initial financed amount; weigh it against the set-up cost",
           font=F_NOTE)

    notes_first = 25
    notes_last = notes_first + len(ORDER) - 1
    _set(ws, "A13", "Execution caveat", font=F_CAVEAT_B, fill=FILL_CAVEAT, wrap=True)
    _merge(ws, "B13:F13",
           f"=INDEX($B${notes_first}:$B${notes_last},MATCH(B10,$A${notes_first}:$A${notes_last},0))",
           font=F_CAVEAT, fill=FILL_CAVEAT, wrap=True)
    ws.row_dimensions[13].height = 58

    _band(ws, 15, "Credit metrics (aircraft-level proxies, linked from each tab)", last_col=6)
    _header(ws, 16, ["Structure", "Initial LTV", "LTV at mid-life", "Asset DSCR",
                     "Weighted average life", "Balloon"])
    cm_first = 17
    for i, key in enumerate(ORDER):
        r = cm_first + i
        _set(ws, f"A{r}", labels[key], font=F_LABEL_B, border=B_ROW)
        for col, field, num in (("B", "ltv0", NUM_PCT1), ("C", "ltv_mid", NUM_PCT1),
                                ("D", "dscr", NUM_X), ("E", "wal", NUM_YRS),
                                ("F", "balloon", NUM_PCT1)):
            _set(ws, f"{col}{r}", link(key, field), font=F_LINK, num=num, align="right",
                 border=B_ROW)
        ws.row_dimensions[r].height = 22
    _merge(ws, "A21:F21",
           "LTV against market value; mid-life value on the Assumptions depreciation curve. "
           "Asset DSCR = the airframe's lease-earning power over the annual financing outflow.",
           font=F_NOTE, wrap=True)
    ws.row_dimensions[21].height = 34

    _band(ws, 23, "Execution notes: what each structure takes to close", last_col=6)
    _header(ws, 24, ["Structure", "What it takes", "", "", "", ""], left=(2,))
    for i, key in enumerate(ORDER):
        r = notes_first + i
        _set(ws, f"A{r}", labels[key], font=F_LABEL_B, border=B_ROW, wrap=True)
        _merge(ws, f"B{r}:F{r}", EXECUTION_NOTES[key], font=F_CALC, wrap=True, border=B_ROW)
        ws.row_dimensions[r].height = 52

    # Highlight the cheapest structure's row in both tables, live.
    for top, bottom in ((first, lastr), (cm_first, cm_first + len(ORDER) - 1),
                        (notes_first, notes_last)):
        ws.conditional_formatting.add(
            f"A{top}:F{bottom}",
            FormulaRule(formula=[f"$A{top}=$B$10"], fill=FILL_LEAD,
                        font=_font(bold=True, color=C_BLUE_DARK)))

    chart = BarChart()
    chart.type = "col"
    chart.title = "Implied all-in annual financing cost"
    chart.legend = None
    chart.add_data(Reference(ws, min_col=2, min_row=3, max_row=lastr), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=first, max_row=lastr))
    # openpyxl leaves axis `delete` unset, which Excel reads as hidden axes.
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.x_axis.title = "Structure"
    chart.y_axis.title = "Implied annual cost"
    chart.y_axis.numFmt = NUM_PCT
    chart.y_axis.scaling.min = 0
    chart.x_axis.tickLblPos = "low"
    chart.y_axis.majorGridlines = ChartLines()
    chart.gapWidth = 70
    chart.series[0].graphicalProperties.solidFill = C_BLUE
    chart.series[0].graphicalProperties.line.solidFill = C_BLUE
    chart.height = 9
    chart.width = 22
    ws.add_chart(chart, "A31")

    return {"first": first, "last": lastr, "cheapest": "$B$10", "edge": "$B$11",
            "worth": "$B$12", "caveat": "$B$13"}


def _build_summary(ws: Worksheet, *, aircraft_label: str, tenor: int, gtf: bool,
                   reference_rate_pct: float, reference_rate_source: str,
                   labels: dict[str, str], cmp: dict, market: dict | None,
                   generated: str) -> None:
    _sheet(ws, {"A": 34, "B": 18, "C": 16, "D": 14, "E": 34, "F": 15}, tab=C_BLUE)
    _title(ws, "Aircraft Financing Structure Comparator",
           f"{aircraft_label}  |  {tenor}-year tenor  |  reference rate {reference_rate_pct:.2f}%"
           f"  |  generated {generated}")

    _merge(ws, "A4:F5",
           "Method: every structure is reduced to the IRR of its incremental cash flows against "
           "simply paying cash for the aircraft at delivery. Because all four share that one "
           "baseline, differences in loan-to-value, fees, tax benefits and the residual given up "
           "all land in a single annual cost, so the four numbers are directly comparable.",
           font=_font(12, color=C_INK), fill=FILL_BOX, wrap=True, border=B_BOX)
    ws.row_dimensions[4].height = 24
    ws.row_dimensions[5].height = 24

    C = "Comparison"
    _band(ws, 7, "Result", last_col=6)
    _set(ws, "A8", "Lowest all-in cost", font=F_LABEL_B)
    _merge(ws, "B8:D8", f"='{C}'!{cmp['cheapest']}", font=F_LINK_KPI)
    _set(ws, "E8", f"='{C}'!{cmp['edge']}", font=F_LINK_B, num=NUM_BPS, align="right")
    _set(ws, "F8", "edge / year", font=F_NOTE)
    ws.row_dimensions[8].height = 30
    _set(ws, "A9", "Execution caveat", font=F_CAVEAT_B, fill=FILL_CAVEAT, wrap=True)
    _merge(ws, "B9:F9", f"='{C}'!{cmp['caveat']}", font=F_CAVEAT, fill=FILL_CAVEAT, wrap=True)
    ws.row_dimensions[9].height = 58

    _band(ws, 11, "Implied all-in annual cost", last_col=6)
    _header(ws, 12, ["Structure", "Implied annual cost", "vs lowest", "Financed",
                     "Accounting (IFRS 16)", "Integrity"], left=(5,))
    for i, _key in enumerate(ORDER):
        r, src = 13 + i, cmp["first"] + i
        _set(ws, f"A{r}", f"='{C}'!A{src}", font=F_LINK_B, border=B_ROW)
        _set(ws, f"B{r}", f"='{C}'!B{src}", font=F_LINK_B, num=NUM_PCT, align="right", border=B_ROW)
        _set(ws, f"C{r}", f"='{C}'!C{src}", font=F_LINK, num=NUM_EDGE, align="right", border=B_ROW)
        _set(ws, f"D{r}", f"='{C}'!D{src}", font=F_LINK, num=NUM_PCT1, align="right", border=B_ROW)
        _set(ws, f"E{r}", f"='{C}'!E{src}", font=F_LINK, align="left", border=B_ROW, indent=1)
        _set(ws, f"F{r}", f"='{C}'!F{src}", font=F_OK, align="right", border=B_ROW)
        ws.row_dimensions[r].height = 22
    ws.conditional_formatting.add(
        "A13:F16", FormulaRule(formula=["$A13=$B$8"], fill=FILL_LEAD,
                               font=_font(bold=True, color=C_BLUE_DARK)))
    ws.conditional_formatting.add(
        "F13:F16", CellIsRule(operator="equal", formula=['"CHECK"'],
                              font=_font(bold=True, color=C_RED)))

    _band(ws, 18, "Deal inputs (edit them on the Assumptions tab)", last_col=6)
    deal = [
        ("Aircraft", R("aircraft"), None),
        ("Delivery price ($m)", R("delivery"), NUM_USD),
        ("Market value today ($m)", R("market_value"), NUM_USD),
        ("Tenor (years)", R("tenor"), NUM_INT),
        ("Reference rate (%)", R("reference_rate_pct"), NUM_RATE),
        ("GTF adjustment (1 = on)", R("gtf"), NUM_INT),
    ]
    for i, (label, ref, num) in enumerate(deal):
        r = 19 + i
        _set(ws, f"A{r}", label, font=F_LABEL, border=B_ROW)
        _set(ws, f"B{r}", f"={ref}", font=F_LINK_B, num=num, align="right" if num else "left",
             border=B_ROW)

    _band(ws, 26, "Data provenance", last_col=6)
    prov = [("Reference rate source", _rate_source_text(reference_rate_source))]
    if market:
        if market.get("sofr_pct") is not None:
            prov.append(("SOFR at export (NY Fed)",
                         f"{market['sofr_pct']:.2f}%, as of {market.get('sofr_as_of') or 'n/a'}"))
        if market.get("eur_str_pct") is not None:
            prov.append(("EUR STR at export (ECB)",
                         f"{market['eur_str_pct']:.2f}%, as of {market.get('eur_str_as_of') or 'n/a'}"))
        if market.get("af_pa_price_eur") is not None:
            prov.append(("AF.PA at export (Yahoo Finance)", f"EUR {market['af_pa_price_eur']:.2f}"))
        if market.get("eurusd") is not None:
            prov.append(("EUR/USD at export (Yahoo Finance)", f"{market['eurusd']:.4f}"))
    prov += [
        ("Sourced assumptions (not live)",
         "Margins, ECA premium, SLL ratchet, JOLCO tax benefit, lease rate factor, depreciation, GTF."),
        ("Deliberately not modelled",
         "Airline credit spreads, bank cost of funds, ECA fee grids, live quotes: not public data."),
    ]
    r = 27
    for label, text in prov:
        _set(ws, f"A{r}", label, font=F_LABEL, border=B_ROW)
        _merge(ws, f"B{r}:F{r}", text, font=F_CALC, border=B_ROW)
        r += 1

    r += 1
    _band(ws, r, "Tabs in this workbook", last_col=6)
    tabs = [
        ("Assumptions", "Every input, with its basis. Blue cells are editable."),
        ("ECA Debt / SLL / JOLCO / SLB", "Key terms, full cash-flow schedule, IRR, integrity checks, credit metrics."),
        ("Comparison", "Side by side, execution notes and the chart."),
    ]
    for tab, text in tabs:
        r += 1
        _set(ws, f"A{r}", tab, font=F_LABEL_B, border=B_ROW)
        _merge(ws, f"B{r}:F{r}", text, font=F_CALC, border=B_ROW)
    r += 2
    _merge(ws, f"A{r}:F{r}",
           "Indicative model for structure comparison. Not a valuation and not investment advice.",
           font=F_NOTE)


# --- public entry point -------------------------------------------------


def build_workbook(*, aircraft_label: str, delivery_price_usd_m: float,
                   market_value_usd_m: float, tenor_years: int,
                   reference_rate_pct: float, reference_rate_source: str,
                   gtf_adjustment: bool, resolved: fm.Assumptions,
                   results: list[fm.StructureResult], cheapest: str,
                   market: dict | None = None) -> bytes:
    REF.clear()
    wb = Workbook()
    _set_default_font(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.properties.creator = "Aircraft Financing Structure Comparator"
    wb.properties.title = f"Financing model - {aircraft_label}"

    ws_summary = wb.active
    ws_summary.title = "Summary"
    ws_assumptions = wb.create_sheet("Assumptions")
    sheets = {key: wb.create_sheet(SHEET[key]) for key in ORDER}
    ws_cmp = wb.create_sheet("Comparison")

    _build_assumptions(
        ws_assumptions, aircraft_label=aircraft_label, delivery=delivery_price_usd_m,
        market_value=market_value_usd_m, tenor=tenor_years,
        reference_rate_pct=reference_rate_pct, reference_rate_source=reference_rate_source,
        gtf=gtf_adjustment, a=resolved,
    )

    builders = {"eca_debt": _eca_sheet, "sll": _sll_sheet, "jolco": _jolco_sheet,
                "slb": _slb_sheet}
    refs = {key: builders[key](sheets[key], tenor_years) for key in ORDER}
    labels = {r.structure: r.label for r in results}

    cmp = _build_comparison(ws_cmp, labels, refs)
    _build_summary(
        ws_summary, aircraft_label=aircraft_label, tenor=tenor_years, gtf=gtf_adjustment,
        reference_rate_pct=reference_rate_pct, reference_rate_source=reference_rate_source,
        labels=labels, cmp=cmp, market=market,
        generated=_dt.datetime.now(_dt.timezone.utc).strftime("%Y-%m-%d %H:%M UTC"),
    )
    wb.active = 0

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
